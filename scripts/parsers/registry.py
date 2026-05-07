from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from scripts.core.ledger import base_row, read_csv, money


SUPPORTED_BANKS = {"abc", "boc", "cmb", "icbc"}
APP_SOURCES = {"alipay", "wechat", "meituan", "douyin"}
REPAYMENT_KEYWORDS = ["repayment", "还款"]
REPAYMENT_ADJUSTMENT_KEYWORDS = [
    "refund",
    "reversal",
    "cashback",
    "discount",
    "coupon",
    "rebate",
    "退款",
    "退回",
    "返还",
    "返现",
    "冲正",
    "撤销",
    "立减",
    "优惠",
    "红包",
]


class IntakeError(RuntimeError):
    pass


@dataclass(frozen=True)
class ParseResult:
    rows: list[dict[str, str]]
    warnings: list[str]


def parse_raw_dir(raw_dir: Path) -> ParseResult:
    root = raw_dir.resolve()
    if not root.exists():
        raise IntakeError(f"Raw directory does not exist: {root}")

    rows: list[dict[str, str]] = []
    warnings: list[str] = []
    known_roots = APP_SOURCES | {"bank"}

    for child in sorted(root.iterdir()):
        if not child.is_dir():
            warnings.append(f"Skipped file at raw root: {child.name}. Put files under raw/<source>/ or raw/bank/<bank>/.")
            continue
        if child.name not in known_roots:
            warnings.append(f"Unknown source folder: {child.relative_to(root)}. Supported: alipay, wechat, meituan, douyin, bank/<boc|cmb|icbc|abc>.")

    for source in sorted(APP_SOURCES):
        source_dir = root / source
        if not source_dir.exists():
            continue
        rows.extend(_parse_app_dir(source, root, source_dir, warnings))

    bank_dir = root / "bank"
    if bank_dir.exists():
        for source_dir in sorted(path for path in bank_dir.iterdir() if path.is_dir()):
            bank_source = source_dir.name.lower()
            if bank_source not in SUPPORTED_BANKS:
                warnings.append(f"Unknown bank folder: {source_dir.relative_to(root)}. Supported banks: abc, boc, cmb, icbc.")
                continue
            rows.extend(_parse_bank_dir(bank_source, root, source_dir, warnings))

    rows.sort(key=lambda row: (row["transaction_time"], row["source"], row["raw_row_number"]))
    return ParseResult(rows=rows, warnings=warnings)


def source_counts(rows: list[dict[str, str]]) -> Counter[str]:
    return Counter(row["source"] for row in rows)


def _parse_app_dir(source: str, root: Path, source_dir: Path, warnings: list[str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in sorted(source_dir.iterdir()):
        suffix = path.suffix.lower()
        if suffix in {".csv", ".xlsx"}:
            rows.extend(_parse_app_csv(source, root, path))
        elif suffix == ".xls":
            warnings.append(f"{path.relative_to(root)} is legacy XLS. Save as CSV or XLSX before import.")
        elif suffix == ".pdf":
            warnings.append(f"{path.relative_to(root)} is PDF. PDF branch is source-specific; use CSV export for stable v1 intake.")
        elif path.is_file():
            warnings.append(f"{path.relative_to(root)} has unsupported extension '{suffix}'.")
    return rows


def _parse_bank_dir(source: str, root: Path, source_dir: Path, warnings: list[str]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for path in sorted(source_dir.iterdir()):
        suffix = path.suffix.lower()
        if suffix in {".csv", ".xlsx"}:
            rows.extend(_parse_bank_csv(source, root, path))
        elif suffix == ".xls":
            warnings.append(f"{path.relative_to(root)} is legacy XLS. Save as CSV or XLSX before import.")
        elif suffix == ".pdf":
            warnings.append(f"{path.relative_to(root)} is PDF. Text-PDF support is planned per bank branch; export CSV when possible.")
        elif path.is_file():
            warnings.append(f"{path.relative_to(root)} has unsupported extension '{suffix}'.")
    return rows


def _parse_app_csv(source: str, root: Path, path: Path) -> list[dict[str, str]]:
    if source == "alipay":
        return _parse_alipay_csv(root, path)
    if source == "wechat":
        return _parse_wechat_csv(root, path)
    if source == "meituan":
        return _parse_meituan_csv(root, path)
    if source == "douyin":
        return _parse_douyin_csv(root, path)
    raise IntakeError(f"Unsupported app source: {source}")


def _parse_alipay_csv(root: Path, path: Path) -> list[dict[str, str]]:
    rows = []
    for idx, row in enumerate(_read_table(path), start=1):
        out = base_row("alipay", root, path, idx, row)
        raw_direction = row.get("raw_direction") or row.get("direction") or ""
        status = row.get("status", "")
        title = _text(row.get("counterparty"), row.get("item_title"), row.get("payment_method"))
        repayment_text = _text(row.get("counterparty"), row.get("item_title"), status)
        if raw_direction == "收入":
            out["direction"] = "income"
            out["normalized_type"] = "refund_in" if _has_any(status + title, ["退款", "refund"]) else "income"
        elif raw_direction in {"支出", "其他", "不计收支"} and _is_confirmed_repayment(repayment_text):
            out["direction"] = "expense"
            out["normalized_type"] = "credit_repayment"
        elif raw_direction == "其他":
            out["direction"] = "neutral"
            out["normalized_type"] = "unknown_wallet_flow"
            out["needs_review"] = "true"
        elif _is_repayment_candidate(repayment_text):
            out["direction"] = "expense"
            out["normalized_type"] = "repayment_candidate_review"
            out["needs_review"] = "true"
        elif _has_any(title, ["top-up", "topup", "充值", "转入"]):
            out["direction"] = "expense"
            out["normalized_type"] = "wallet_topup"
        else:
            out["direction"] = "expense"
            out["normalized_type"] = "merchant_payment"
        out["normalized_status"] = "refund" if _has_any(status, ["退款", "refund"]) else "success"
        rows.append(out)
    return rows


def _parse_wechat_csv(root: Path, path: Path) -> list[dict[str, str]]:
    rows = []
    for idx, row in enumerate(_read_table(path), start=1):
        out = base_row("wechat", root, path, idx, row)
        flow = row.get("income_expense") or row.get("direction") or ""
        title = _text(row.get("counterparty"), row.get("item_title"), row.get("payment_method"))
        if flow == "收入":
            out["direction"] = "income"
            out["normalized_type"] = "refund_in" if _has_any(title, ["退款", "refund"]) else "income"
            out["normalized_status"] = "refund" if out["normalized_type"] == "refund_in" else "success"
        elif _is_confirmed_repayment(title):
            out["direction"] = "expense"
            out["normalized_type"] = "personal_repayment"
            out["normalized_status"] = "success"
        elif _is_repayment_candidate(title):
            out["direction"] = "expense"
            out["normalized_type"] = "repayment_candidate_review"
            out["normalized_status"] = "success"
            out["needs_review"] = "true"
        else:
            out["direction"] = "expense"
            out["normalized_type"] = "merchant_payment"
            out["normalized_status"] = "success"
        rows.append(out)
    return rows


def _parse_meituan_csv(root: Path, path: Path) -> list[dict[str, str]]:
    rows = []
    for idx, row in enumerate(_read_table(path), start=1):
        out = base_row("meituan", root, path, idx, row)
        kind = row.get("kind", "")
        title = _text(kind, row.get("item_title"), row.get("counterparty"))
        if _has_any(title, ["退款", "refund"]):
            out["direction"] = "income"
            out["normalized_type"] = "refund_in"
        elif _is_confirmed_repayment(title):
            out["direction"] = "expense"
            out["normalized_type"] = "credit_repayment"
        elif _is_repayment_candidate(title):
            out["direction"] = "expense"
            out["normalized_type"] = "repayment_candidate_review"
            out["needs_review"] = "true"
        else:
            out["direction"] = "expense"
            out["normalized_type"] = "merchant_payment"
        out["normalized_status"] = "refund" if out["normalized_type"] == "refund_in" else "success"
        rows.append(out)
    return rows


def _parse_douyin_csv(root: Path, path: Path) -> list[dict[str, str]]:
    rows = []
    for idx, row in enumerate(_read_table(path), start=1):
        out = base_row("douyin", root, path, idx, row)
        flow = row.get("direction") or row.get("income_expense") or ""
        title = _text(row.get("payment_method"), row.get("counterparty"), row.get("item_title"), row.get("kind"))
        if flow == "收入":
            out["direction"] = "income"
            out["normalized_type"] = "platform_reward" if _has_any(title, ["reward", "campaign", "奖励", "补贴"]) else "refund_in"
            out["normalized_status"] = "success" if out["normalized_type"] == "platform_reward" else "refund"
        elif _is_confirmed_repayment(title):
            out["direction"] = "expense"
            out["normalized_type"] = "credit_repayment"
            out["normalized_status"] = "success"
        elif _is_repayment_candidate(title):
            out["direction"] = "expense"
            out["normalized_type"] = "repayment_candidate_review"
            out["normalized_status"] = "success"
            out["needs_review"] = "true"
        else:
            out["direction"] = "expense"
            out["normalized_type"] = "merchant_payment"
            out["normalized_status"] = "success"
        rows.append(out)
    return rows


def _parse_bank_csv(source: str, root: Path, path: Path) -> list[dict[str, str]]:
    rows = []
    for idx, row in enumerate(_read_table(path), start=1):
        out = base_row(source, root, path, idx, row)
        credit = row.get("credit", "")
        debit = row.get("debit", "")
        desc = _text(row.get("description"), row.get("item_title"), row.get("counterparty"))
        if credit:
            out["direction"] = "income"
            out["amount"] = f"{money(credit):.2f}"
            out["normalized_type"] = "salary" if _has_any(desc, ["salary", "工资", "薪资"]) else "bank_income"
        elif debit:
            out["direction"] = "expense"
            out["amount"] = f"{money(debit):.2f}"
            out["normalized_type"] = "internal_transfer" if _has_any(desc, ["transfer", "转账", "充值", "top-up"]) else "bank_payment"
        else:
            out["direction"] = "neutral"
            out["amount"] = "0.00"
            out["normalized_type"] = "unknown_wallet_flow"
            out["needs_review"] = "true"
        out["normalized_status"] = "success"
        rows.append(out)
    return rows


def _text(*values: str | None) -> str:
    return " ".join(value for value in values if value)


def _has_any(value: str, needles: list[str]) -> bool:
    lowered = value.lower()
    return any(needle.lower() in lowered for needle in needles)


def _is_repayment_candidate(value: str) -> bool:
    return _has_any(value, REPAYMENT_KEYWORDS)


def _is_confirmed_repayment(value: str) -> bool:
    return _is_repayment_candidate(value) and not _has_any(value, REPAYMENT_ADJUSTMENT_KEYWORDS)


def _read_table(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".csv":
        return read_csv(path)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise IntakeError("XLSX/XLS intake requires openpyxl. Install it locally or export CSV.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    output: list[dict[str, str]] = []
    for values in rows[1:]:
        record = {}
        for header, value in zip(headers, values):
            if header:
                record[header] = "" if value is None else str(value)
        if any(record.values()):
            output.append(record)
    return output
